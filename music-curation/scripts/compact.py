import openpyxl,sys
wb=openpyxl.load_workbook(sys.argv[1]); ws=wb['Sangeet Dance Tracks']
ok=lambda v: isinstance(v,str) and v.strip()
kept=set(); lens=[]
for col in range(2,10):
    head=ws.cell(row=2,column=col).value
    vals=[ws.cell(row=r,column=col).value.strip() for r in range(3,ws.max_row+1) if ok(ws.cell(row=r,column=col).value)]
    for r in range(3,ws.max_row+1): ws.cell(row=r,column=col).value=None
    for i,v in enumerate(vals): ws.cell(row=3+i,column=col).value=v
    kept|={(head,v) for v in vals}; lens.append((head,len(vals)))
n=max(l for _,l in lens)
if ws.max_row>n+2: ws.delete_rows(n+3,ws.max_row-(n+2))
# keep the Track Details tab in sync with songs still on the main sheet
d=wb['Track Details']
rows=[[c.value for c in r] for r in d.iter_rows(min_row=2,max_row=d.max_row)]
rows=[r for r in rows if (r[0],r[1]) in kept]
for r in range(2,d.max_row+1):
    for c in range(1,11): d.cell(row=r,column=c).value=None
for i,r in enumerate(rows):
    for j,v in enumerate(r): d.cell(row=2+i,column=1+j).value=v
if d.max_row>len(rows)+1: d.delete_rows(len(rows)+2,d.max_row-len(rows)-1)
d.auto_filter.ref=f'A1:J{len(rows)+1}'
wb.save(sys.argv[2]); print(lens, 'details rows',len(rows),'kept',len(kept))
