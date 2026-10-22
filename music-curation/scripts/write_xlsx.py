import json,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Border,Side,Alignment
from openpyxl.utils import get_column_letter
F=json.load(open('final.json'))
COLS=[('English','English Pop'),('OldHindi','Old Hindi Pop 80s-00s'),('NewHindi','New Hindi Pop 00s-20s'),('Punjabi','Punjabi'),
      ('Tamil','Tamil'),('Telugu','Telugu'),('Kannada','Kannada'),('Malayalam','Malayalam')]
thin=Side(style='thin',color='FF000000'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
white=PatternFill('solid',fgColor='FFFFFFFF'); yellow=PatternFill('solid',fgColor='FFFFF2CC')
hfont=Font(name='Arial',bold=True); bfont=Font(name='Arial',color='FF000000')
wb=Workbook(); ws=wb.active; ws.title='Sangeet Dance Tracks'
ws.column_dimensions['A'].width=5
nmax=max(len(F[c]) for c,_ in COLS)
for i,(c,h) in enumerate(COLS):
    col=i+2; L=get_column_letter(col)
    ws.column_dimensions[L].width=max(30, min(48, max(len(r[0]) for r in F[c])*0.95))
    hc=ws.cell(row=2,column=col,value=h); hc.font=hfont; hc.fill=yellow; hc.border=B
    for j in range(nmax):
        cell=ws.cell(row=3+j,column=col); cell.font=bfont; cell.fill=white; cell.border=B
        cell.alignment=Alignment(vertical='top')
        if j<len(F[c]): cell.value=F[c][j][0]
ws.freeze_panes='B3'
# Reference tab: which file/version in the Traktor collection each song maps to
d=wb.create_sheet('Track Details')
hdr=['Category','Song (as listed)','Title in Collection','Artist','Version / Mix','Traktor Genre','BPM','Play Count','Other Versions','File Name']
for k,h in enumerate(hdr,1):
    x=d.cell(row=1,column=k,value=h); x.font=hfont; x.fill=yellow; x.border=B
r=2
for c,h in COLS:
    for name,vs in F[c]:
        uniq={}
        for t in vs: uniq.setdefault(t['title'].strip().lower(),t)
        t=max(vs,key=lambda t:(t['plays'],len(t['playlists'])))
        vals=[h,name,t['title'],t['artist'],t['mix'],t['genre'],t['bpm'] or None,t['plays'],len(uniq)-1,t['file']]
        for k,v in enumerate(vals,1):
            x=d.cell(row=r,column=k,value=v); x.font=bfont; x.border=B
        r+=1
for k,w in enumerate([22,36,40,30,20,26,8,11,14,60],1): d.column_dimensions[get_column_letter(k)].width=w
d.freeze_panes='C2'; d.auto_filter.ref=f'A1:J{r-1}'
wb.save(sys.argv[1]); print('rows',r-2,'max col len',nmax)
