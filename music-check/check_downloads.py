#!/usr/bin/env python3
"""
Check which songs in the Shazam master Excel are already downloaded on this Mac.

It scans your music folders for audio files and compares each file name (and,
when available, its title/artist tags) against every song in the master sheet.
Matching ignores capitals, punctuation and spacing, and treats common
Hindi/English spelling variants as the same word, e.g.
"Yeh Dosti" = "Ye Dosti" = "YehDosti", "Deewana" = "Diwana", "Pyaar" = "Pyar".

Two columns are added to the master sheet (a new copy is written; your
original file is never changed):

    Available on Mac?   Yes / No / Check
    Download Comments   which file matched, or the similar file(s) to verify

"Check" means a similar-looking file was found and you should decide by hand
whether it is the same track.

Usage:
    python3 check_downloads.py MASTER.xlsx [FOLDER ...]
    python3 check_downloads.py MASTER.xlsx --list music_files.txt

With no FOLDER, your whole home folder is scanned (skipping ~/Library and
hidden folders) plus any external drives under /Volumes.
"""
import argparse
import copy
import datetime
import os
import re
import sys
import unicodedata
from collections import Counter
from difflib import SequenceMatcher

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is missing. Run ./run.sh instead (it installs it), "
             "or: python3 -m pip install openpyxl")

try:                       # optional: reads title/artist tags inside the files
    import mutagen
except ImportError:
    mutagen = None

try:                       # optional: faster fuzzy matching
    from rapidfuzz import fuzz

    def similarity(a, b):
        return fuzz.ratio(a, b) / 100.0
except ImportError:
    def similarity(a, b):
        return SequenceMatcher(None, a, b).ratio()

AUDIO_EXTS = {".mp3", ".m4a", ".aac", ".wav", ".flac", ".aif", ".aiff", ".alac",
              ".ogg", ".opus", ".wma", ".m4p", ".mp4", ".caf"}
SKIP_DIRS = {"Library", "node_modules", "Photos Library.photoslibrary", ".Trash"}

# Words that say nothing about which song a file is.
NOISE = {"official", "audio", "video", "lyrics", "lyric", "lyrical", "hd", "hq",
         "4k", "320kbps", "kbps", "128kbps", "mp3", "full", "song", "songs",
         "version", "from", "feat", "ft", "featuring", "with", "the", "a", "an",
         "and", "music", "visualizer", "movie", "film", "remastered", "copy"}

SHEET = "Shazam Master"
COL_AVAIL = "Available on Mac?"
COL_NOTE = "Download Comments"


# --- Text normalisation ------------------------------------------------------

def plain(text):
    """Lowercase, strip accents, turn punctuation into spaces."""
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    text = text.replace("'", "").replace("’", "").replace("&", " ")
    return re.sub(r"[^\w]+", " ", text).replace("_", " ").strip()


def sound_key(word):
    """Collapse spelling variants of the same sounding word (Hinglish-aware)."""
    w = word
    for a, b in (("ee", "i"), ("oo", "u"), ("ph", "f"), ("ck", "k"), ("q", "k"),
                 ("w", "v"), ("z", "j")):
        w = w.replace(a, b)
    w = re.sub(r"([bcdgjkpst])h", r"\1", w)          # bh dh kh gh sh th ch -> b d k g s t c
    w = re.sub(r"([aeiou])h(?=[^aeiou]|$)", r"\1", w)  # yeh -> ye, rehna -> rena
    w = re.sub(r"(.)\1+", r"\1", w)                   # pyaar -> pyar, dostii -> dosti
    w = re.sub(r"(?<=.)(ie|y)$", "i", w)              # kabhie -> kabi, dosty -> dosti
    return w


def tokens(text, drop_noise=True):
    words = plain(text).split()
    if drop_noise:
        words = [w for w in words if w not in NOISE]
    words = [re.sub(r"^\d{1,3}$", "", w) if i == 0 and len(words) > 1 else w
             for i, w in enumerate(words)]            # leading track number "01"
    return [sound_key(w) for w in words if w]


def core_title(title):
    """Song title without (brackets), ' - subtitle' and feat. credits."""
    t = re.sub(r"[\(\[\{].*?[\)\]\}]", " ", str(title))
    t = re.split(r"\s[-–—|]\s", t)[0]
    t = re.split(r"\b(?:feat|ft|featuring)\b\.?", t, flags=re.I)[0]
    return t


def artist_names(artist):
    """'A, B & C feat. D' -> token lists for each individual artist."""
    parts = re.split(r",|&|\band\b|\bx\b|\bfeat\.?|\bft\.?|\bwith\b", str(artist or ""), flags=re.I)
    names = []
    for p in parts:
        t = [w for w in tokens(p) if len(w) >= 3]
        if t:
            names.append(t)
    return names


def trigrams(text):
    return {text[i:i + 3] for i in range(len(text) - 2)} or ({text} if text else set())


def word_match(a, b):
    if a == b:
        return True
    return len(a) >= 5 and len(b) >= 5 and similarity(a, b) >= 0.8


# --- Collect files -----------------------------------------------------------

def walk_audio(roots):
    for root in roots:
        root = os.path.expanduser(root)
        if not os.path.isdir(root):
            print(f"  (skipping, not a folder: {root})")
            continue
        print(f"  scanning {root} ...")
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames
                           if not d.startswith(".") and d not in SKIP_DIRS]
            for name in filenames:
                if os.path.splitext(name)[1].lower() in AUDIO_EXTS and not name.startswith("._"):
                    yield os.path.join(dirpath, name)


def read_tags(path):
    if mutagen is None:
        return None
    try:
        f = mutagen.File(path, easy=True)
    except Exception:
        return None
    if not f or not f.tags:
        return None
    title = (f.tags.get("title") or [""])[0]
    artist = (f.tags.get("artist") or [""])[0]
    return f"{artist} - {title}".strip(" -") or None


class AudioFile:
    def __init__(self, path, tag_name=None):
        self.path = path
        self.name = os.path.basename(path)
        stem = os.path.splitext(self.name)[0]
        self.variants = []
        for text in filter(None, (stem, tag_name)):
            toks = tokens(text)
            self.variants.append((toks, "".join(toks)))


def short_path(path):
    home = os.path.expanduser("~")
    return "~" + path[len(home):] if path.startswith(home) else path


# --- Matching ----------------------------------------------------------------

def score(title_toks, title_joined, artists, file_toks):
    """Return (title coverage 0-1, all exact?, artist found?, extra words)."""
    if not title_toks:
        return 0.0, False, False, 0
    used = set()
    hits = exact = 0
    for t in title_toks:
        for i, f in enumerate(file_toks):
            if i not in used and word_match(t, f):
                used.add(i)
                hits += 1
                exact += t == f
                break
    coverage = hits / len(title_toks)
    all_exact = exact == len(title_toks)
    if coverage < 1 and len(title_joined) >= 4:
        # Same words, different spacing: "Yeh Dosti" vs "YehDosti",
        # "Chor Bazaare" vs "Chorbazari". Compare against runs of whole words.
        for i in range(len(file_toks)):
            run = ""
            for j in range(i, min(len(file_toks), i + len(title_toks) + 1)):
                run += file_toks[j]
                if run == title_joined:
                    coverage, all_exact, used = 1.0, True, set(range(i, j + 1))
                elif (len(title_joined) >= 5 and abs(len(run) - len(title_joined)) <= 1
                      and similarity(run, title_joined) >= 0.85 and coverage < 0.85):
                    coverage, used = 0.85, set(range(i, j + 1))
    artist_hit = False
    for name in artists:
        idx = [i for i, f in enumerate(file_toks) if any(word_match(w, f) for w in name)]
        if len(idx) >= min(len(name), 2) or (len(name) == 1 and idx):
            artist_hit = True
            used.update(idx)
    extra = len([i for i in range(len(file_toks)) if i not in used])
    return coverage, all_exact, artist_hit, extra


def classify(title, artist, files, index):
    title_toks = tokens(core_title(title)) or tokens(title)
    if not title_toks or plain(title) in ("unknown", ""):
        return "Check", "Song name is unknown in the Shazam data, so it could not be checked"
    title_joined = "".join(title_toks)
    artists = artist_names(artist)
    distinctive = len(title_toks) >= 3 or len(title_joined) >= 12

    candidates = set()
    for t in title_toks:
        if len(t) >= 3:
            candidates |= index["words"].get(t, set())
    grams = trigrams(title_joined)
    if grams:
        shared = Counter(i for g in grams for i in index["grams"].get(g, ()))
        need = max(2, 0.6 * len(grams))
        candidates |= {i for i, n in shared.items() if n >= need}

    found, similar = [], []
    for i in candidates:
        f = files[i]
        best = None
        for ftoks, _ in f.variants:
            cov, all_exact, a_hit, extra = score(title_toks, title_joined, artists, ftoks)
            rank = cov * 100 + a_hit * 25 + all_exact * 5 - extra
            if best is None or rank > best[0]:
                best = (rank, cov, all_exact, a_hit, extra)
        rank, cov, all_exact, a_hit, extra = best
        if cov == 1 and all_exact and (a_hit or (distinctive and extra <= 3)):
            found.append((rank, f))
        elif ((cov == 1 and (len(title_toks) >= 2 or a_hit or extra <= 1))
              or (cov >= 0.75 and len(title_toks) >= 2) or (a_hit and cov >= 0.5)):
            similar.append((rank, f))

    found.sort(key=lambda x: -x[0])
    similar.sort(key=lambda x: -x[0])
    if found:
        names = "; ".join(f"{f.name}  [{short_path(os.path.dirname(f.path))}]" for _, f in found[:3])
        more = f" (+{len(found) - 3} more copies)" if len(found) > 3 else ""
        note = f"Found: {names}{more}"
        return "Yes", note
    if similar:
        names = "; ".join(f"{f.name}  [{short_path(os.path.dirname(f.path))}]" for _, f in similar[:3])
        return "Check", f"Similar track found - please verify: {names}"
    return "No", ""


# --- Excel -------------------------------------------------------------------

def write_results(master_path, out_path, results, scanned, roots):
    wb = openpyxl.load_workbook(master_path)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    for name in (COL_AVAIL, COL_NOTE):
        if name not in header:
            header.append(name)
            ws.cell(row=1, column=len(header), value=name)
    ca, cn = header.index(COL_AVAIL) + 1, header.index(COL_NOTE) + 1
    hdr_src, body_src = ws.cell(row=1, column=1), ws.cell(row=2, column=1)
    fills = {"Yes": "E2EFDA", "Check": "FFF2CC", "No": None}
    for col in (ca, cn):
        h = ws.cell(row=1, column=col)
        h.font, h.fill = copy.copy(hdr_src.font), copy.copy(hdr_src.fill)
        h.border, h.alignment = copy.copy(hdr_src.border), copy.copy(hdr_src.alignment)
    for row, (avail, note) in results.items():
        for col, val in ((ca, avail), (cn, note)):
            c = ws.cell(row=row, column=col, value=val)
            c.font, c.border = copy.copy(body_src.font), copy.copy(body_src.border)
            c.alignment = copy.copy(body_src.alignment)
            c.fill = PatternFill("solid", fgColor=fills[avail]) if fills[avail] else PatternFill()
    ws.column_dimensions[get_column_letter(ca)].width = 14
    ws.column_dimensions[get_column_letter(cn)].width = 70
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"

    if "Download Check" in wb.sheetnames:
        del wb["Download Check"]
    s = wb.create_sheet("Download Check")
    bold, body = Font(name="Arial", size=11, bold=True), Font(name="Arial", size=10)
    col = get_column_letter(ca)
    rows = [
        ("Status", "Songs", "Meaning"),
        ("Yes", f"=COUNTIF('{SHEET}'!${col}:${col},A2)", "A file with this song name (and usually the artist) was found"),
        ("Check", f"=COUNTIF('{SHEET}'!${col}:${col},A3)", "A similar file was found - verify by hand using Download Comments"),
        ("No", f"=COUNTIF('{SHEET}'!${col}:${col},A4)", "Nothing similar found - still to download"),
        ("Total", "=SUM(B2:B4)", ""),
        ("", "", ""),
        ("Checked on", datetime.date.today().isoformat(), ""),
        ("Audio files scanned", scanned, ""),
        ("Folders scanned", "; ".join(roots), ""),
    ]
    thin = Side(style="thin")
    for r, vals in enumerate(rows, 1):
        for c, v in enumerate(vals, 1):
            cell = s.cell(row=r, column=c, value=v)
            cell.font = bold if r == 1 or c == 1 else body
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r <= 5:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if r == 1:
                    cell.fill = PatternFill("solid", fgColor="C6E0B4")
    s.column_dimensions["A"].width, s.column_dimensions["B"].width = 22, 14
    s.column_dimensions["C"].width = 70
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("master", help="the Shazam master .xlsx")
    ap.add_argument("folders", nargs="*", help="folders to scan (default: home folder + /Volumes)")
    ap.add_argument("--list", help="text file with one audio file path per line, instead of scanning")
    ap.add_argument("--out", help="output .xlsx (default: <master>_with_downloads.xlsx)")
    args = ap.parse_args()

    out = args.out or re.sub(r"\.xlsx$", "", args.master, flags=re.I) + "_with_downloads.xlsx"
    if args.list:
        with open(args.list, encoding="utf-8", errors="replace") as fh:
            paths = [ln.strip() for ln in fh
                     if os.path.splitext(ln.strip())[1].lower() in AUDIO_EXTS]
        roots = [f"file list: {os.path.basename(args.list)}"]
        tags = False
    else:
        roots = args.folders or (["~"] + [os.path.join("/Volumes", v) for v in sorted(os.listdir("/Volumes"))
                                         if not v.startswith("Macintosh HD")] if os.path.isdir("/Volumes") else ["~"])
        print("Looking for audio files:")
        paths = list(walk_audio(roots))
        tags = mutagen is not None
    print(f"Found {len(paths)} audio files.")
    if tags:
        print("Reading title/artist tags inside the files ...")

    files = [AudioFile(p, read_tags(p) if tags else None) for p in paths]
    index = {"words": {}, "grams": {}}
    for i, f in enumerate(files):
        for toks, joined in f.variants:
            for t in toks:
                index["words"].setdefault(t, set()).add(i)
            for g in trigrams(joined):
                index["grams"].setdefault(g, set()).add(i)

    wb = openpyxl.load_workbook(args.master, read_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"Sheet '{SHEET}' not found in {args.master}")
    ws = wb[SHEET]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ti, ai = header.index("Song Name"), header.index("Artist")
    results = {}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[ti] is None and row[ai] is None:
            continue
        results[r] = classify(row[ti], row[ai], files, index)
    wb.close()

    write_results(args.master, out, results, len(paths), roots)
    counts = {k: sum(1 for v in results.values() if v[0] == k) for k in ("Yes", "Check", "No")}
    print(f"\nDone. Yes: {counts['Yes']}   Check: {counts['Check']}   No: {counts['No']}")
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
